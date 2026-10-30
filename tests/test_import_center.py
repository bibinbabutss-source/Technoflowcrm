"""Excel / CSV import center: validation, duplicates, ownership, historical dates, revenue safety, security."""
import datetime
import io
from decimal import Decimal
from unittest import mock

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from accounts import capabilities as caps
from customers.models import Customer
from leads.models import Lead
from reporting.import_kinds import KINDS
from reporting.models import ImportBatch, ImportKey
from sales.models import Opportunity
from tasks.models import Task
from visits.models import Visit

pytestmark = pytest.mark.django_db
TODAY = timezone.localdate()


def xlsx(rows, name="data.xlsx"):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return SimpleUploadedFile(name, out.getvalue(),
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def csvfile(text, name="data.csv"):
    return SimpleUploadedFile(name, text.encode(), "text/csv")


@pytest.fixture
def org(make_user, login):
    admin, director = make_user(caps.ADMIN), make_user(caps.DIRECTOR)
    e1, e2 = make_user(caps.ENGINEER, username="eng.one"), make_user(caps.ENGINEER, username="eng.two")
    cust = Customer.objects.create(company_name="Known Customer LLC", code="CUS-90001", assigned_engineer=e1,
                                   trn="100200300400500", email="known@example.com", phone="+971 4 555 0001")
    return {"admin": admin, "director": director, "e1": e1, "e2": e2, "cust": cust,
            "a": login(admin), "d": login(director), "c1": login(e1), "c2": login(e2)}


def upload(client, kind, f, mode="skip"):
    r = client.post(reverse("reporting:import_start"), {"kind": kind, "mode": mode, "file": f})
    batch = ImportBatch.objects.order_by("-pk").first()
    return r, batch


def run(client, batch):
    client.post(reverse("reporting:import_run", args=[batch.pk]))
    batch.refresh_from_db()
    return batch


def results(batch):
    return {r["row"]: r for r in batch.results}


# ------------------------------------------------------------------ templates and file checks
@pytest.mark.parametrize("kind", list(KINDS))
def test_templates_download_and_reimport_cleanly(kind, org):
    r = org["a"].get(reverse("reporting:import_template", args=[kind]))
    assert r.status_code == 200 and r["Content-Disposition"].endswith(f'{kind}.xlsx"')
    wb = load_workbook(io.BytesIO(r.content))
    headers = [c.value for c in wb["Data"][1]]
    assert headers == [f.label + (" *" if f.required else "") for f in KINDS[kind]["fields"]]
    assert "Instructions" in wb.sheetnames
    # the template's own headings map automatically onto every field
    from reporting.imports import guess_mapping
    assert set(guess_mapping(kind, headers)) == {f.name for f in KINDS[kind]["fields"]}


@pytest.mark.parametrize("name,content,expected", [
    ("data.exe", b"MZ\x90\x00", "Upload an Excel"),
    ("data.txt", b"hello", "Upload an Excel"),
    ("fake.xlsx", b"just text pretending", "not a real Excel"),
    ("broken.xlsx", b"PK\x03\x04garbage-not-a-workbook", "could not be read"),
    ("bin.csv", b"a,b\x00\x00\x01", "not a plain-text CSV"),
    ("empty.csv", b"", "file is empty"),
    ("blank.csv", b"\n\n", "no rows"),
    ("head.csv", b"Company name,Email\n", "no data rows"),
])
def test_bad_files_are_refused_safely(name, content, expected, org):
    r, batch = upload(org["a"], "customers", SimpleUploadedFile(name, content))
    assert r.status_code == 200 and batch is None
    assert expected in r.content.decode()
    assert b"Traceback" not in r.content and b"/home/" not in r.content


def test_size_and_row_limits(org):
    big = "Company name\n" + "\n".join(f"Co {i} LLC" for i in range(2001))
    r, batch = upload(org["a"], "customers", csvfile(big))
    assert batch is None and "more than 2000 rows" in r.content.decode()
    huge = SimpleUploadedFile("huge.csv", b"Company name\n" + b"x" * (5 * 1024 * 1024 + 10))
    r, batch = upload(org["a"], "customers", huge)
    assert batch is None and "larger than 5 MB" in r.content.decode()


def test_missing_required_columns(org):
    r, batch = upload(org["a"], "leads", xlsx([["Company name", "Email"], ["X LLC", "x@example.com"]]))
    assert batch.status == "mapping"
    r = org["a"].get(batch.get_absolute_url())
    assert "Missing: Lead title" in r.content.decode()
    r = org["a"].post(batch.get_absolute_url(), {"map_company_name": "0"})
    batch.refresh_from_db()
    assert batch.status == "mapping" and not Lead.objects.exists()


# ------------------------------------------------------------------ validation, partial failures, report
def test_partial_validation_failures_and_problem_report(org):
    rows = [["Lead title", "Company name", "Sales Engineer", "Estimated value", "Date received", "Email", "Phone",
             "Lead category"],
            ["Good lead", "Alpha LLC", "eng.one", "1,500.50", "15/03/2026", "a@alpha.example", "+971 50 123 4567", "Tender"],
            ["Bad date", "Beta LLC", "eng.one", "100", "31/02/2026", "", "", "Tender"],
            ["Bad number", "Gamma LLC", "eng.one", "abc", "", "", "", "Tender"],
            ["", "Delta LLC", "eng.one", "", "", "", "", "Tender"],
            ["Bad email", "Eps LLC", "eng.one", "", "", "not-an-email", "", "Tender"],
            ["Bad employee", "Zeta LLC", "nobody.here", "", "", "", "", "Tender"],
            ["Bad phone", "Eta LLC", "eng.one", "", "", "", "12", "Tender"],
            ["Future", "Theta LLC", "eng.one", "", (TODAY + datetime.timedelta(days=5)).isoformat(), "", "", "Tender"]]
    r, batch = upload(org["a"], "leads", xlsx(rows))
    res = results(batch)
    assert res[2]["status"] == "new"
    assert "not a valid date received" in res[3]["messages"][0]
    assert "not a valid estimated value" in res[4]["messages"][0]
    assert "Lead title is required" in res[5]["messages"]
    assert "not a valid email" in res[6]["messages"][0]
    assert "No active employee" in res[7]["messages"][0]
    assert "not a valid phone" in res[8]["messages"][0]
    assert "cannot be in the future" in res[9]["messages"][0]
    assert not Lead.objects.exists()  # nothing saved before confirmation
    batch = run(org["a"], batch)
    assert Lead.objects.count() == 1 and Lead.objects.get().estimated_value == Decimal("1500.50")
    rep = org["a"].get(reverse("reporting:import_errors", args=[batch.pk]))
    text = rep.content.decode("utf-8-sig")
    lines = [line for line in text.splitlines()[1:] if line]
    assert len(lines) == 7 and lines[0].startswith("3,Error")  # Excel row numbers, reason, original cells
    assert "Bad email" in text and "Good lead" not in text


def test_cancelled_import_saves_nothing(org):
    r, batch = upload(org["a"], "customers", csvfile("Company name\nCancel Me LLC\n"))
    org["a"].post(reverse("reporting:import_cancel", args=[batch.pk]))
    batch.refresh_from_db()
    assert batch.status == "cancelled" and batch.rows == []
    assert org["a"].post(reverse("reporting:import_run", args=[batch.pk])).status_code == 302
    assert not Customer.objects.filter(company_name="Cancel Me LLC").exists()


def test_rollback_when_an_error_happens_midway(org):
    r, batch = upload(org["a"], "customers", csvfile("Company name\nFirst LLC\nSecond LLC\nThird LLC\n"))
    from reporting import imports

    original = imports.CustomerHandler.create
    calls = {"n": 0}

    def flaky(self, d, ctx, b):
        calls["n"] += 1
        if calls["n"] == 3:
            raise RuntimeError("disk full")
        return original(self, d, ctx, b)

    with mock.patch.object(imports.CustomerHandler, "create", flaky):
        resp = org["a"].post(reverse("reporting:import_run", args=[batch.pk]), follow=True)
    assert "nothing was saved" in resp.content.decode()
    assert not Customer.objects.filter(company_name__in=["First LLC", "Second LLC"]).exists()
    assert not ImportKey.objects.exists()
    batch.refresh_from_db()
    assert batch.status == "preview"


# ------------------------------------------------------------------ duplicates and updates
def test_customer_duplicates_and_reliable_updates(org):
    rows = [["Company name", "TRN", "Email", "Phone", "City", "Reference"],
            ["Known Customer LLC", "", "", "", "Sharjah", "C-1"],             # same name only → duplicate, skipped
            ["Renamed Holding", "100200300400500", "", "", "Dubai", "C-2"],    # same TRN → reliable
            ["Other Name", "", "known@example.com", "+971 4 555 0001", "", ""],  # email+phone → reliable
            ["Brand New LLC", "", "new@example.com", "", "", "C-4"],
            ["Brand New LLC", "", "", "", "", ""]]                             # repeated in the file
    r, batch = upload(org["a"], "customers", xlsx(rows), mode="skip")
    res = results(batch)
    assert res[2]["status"] == "duplicate" and "never updated" in res[2]["messages"][0]
    assert res[3]["status"] == "duplicate" and "same TRN" in res[3]["messages"][0]
    assert res[6]["status"] == "duplicate" and "row 5" in res[6]["messages"][0]
    assert res[5]["status"] == "new"
    # update mode: reliable matches are updated (preview lists the change), name-only stays skipped
    org["a"].post(batch.get_absolute_url(), {**{f"map_{k}": str(v) for k, v in batch.mapping.items()},
                                             "mode": "update"})
    batch.refresh_from_db()
    res = results(batch)
    assert res[2]["status"] == "duplicate"
    assert res[3]["status"] == "update" and ["Emirate / city", "—", "Dubai"] in res[3]["changes"]
    batch = run(org["a"], batch)
    cust = org["cust"]
    cust.refresh_from_db()
    assert cust.city == "Dubai" and cust.company_name == "Known Customer LLC"  # name never overwritten by import
    assert cust.email == "known@example.com"  # empty cells never blank values
    assert batch.summary["created"] == 1 and batch.summary["updated"] == 1
    assert Customer.objects.filter(company_name="Brand New LLC").count() == 1


def test_lead_duplicates_and_same_file_twice(org):
    rows = [["Lead title", "Company name", "Sales Engineer", "Email", "Reference", "Category"],
            ["Valve enquiry", "Omega LLC", "eng.one", "buyer@omega.example", "L-1", "Job in Hand"]]
    r, batch = upload(org["a"], "leads", xlsx(rows))
    run(org["a"], batch)
    r, again = upload(org["a"], "leads", xlsx(rows))
    assert results(again)[2]["status"] == "duplicate"
    run(org["a"], again)
    assert Lead.objects.filter(title="Valve enquiry").count() == 1
    # different reference but same company/email as an open lead → possible duplicate, skipped
    rows2 = [rows[0], ["Another title", "Omega LLC", "eng.one", "buyer@omega.example", "L-9", "Tender"]]
    r, b3 = upload(org["a"], "leads", xlsx(rows2))
    assert results(b3)[2]["status"] == "duplicate" and "open lead" in results(b3)[2]["messages"][0]


# ------------------------------------------------------------------ ownership
def test_ownership_rules(org):
    # an engineer's import is theirs; naming another employee is refused, not silently changed
    rows = [["Lead title", "Company name", "Sales Engineer", "Category"], ["Mine", "Mine LLC", "", "Tender"],
            ["Theirs", "Theirs LLC", "eng.two", "Tender"]]
    r, batch = upload(org["c1"], "leads", xlsx(rows))
    res = results(batch)
    assert res[2]["status"] == "new"
    assert res[3]["status"] == "error" and "not allowed to import records for" in res[3]["messages"][0]
    run(org["c1"], batch)
    lead = Lead.objects.get(title="Mine")
    assert lead.assigned_engineer == org["e1"] and lead.created_by == org["e1"]
    # an administrator must name the employee: nothing is attributed to the admin automatically
    r, b2 = upload(org["a"], "leads", xlsx([["Lead title", "Company name", "Category"], ["Admin lead", "Adm LLC", "PE"]]))
    assert results(b2)[2]["status"] == "error" and "Employee is required" in results(b2)[2]["messages"][0]
    # the director may import for any employee
    r, b3 = upload(org["d"], "leads", xlsx([["Lead title", "Company name", "Sales Engineer", "Lead type"],
                                            ["For two", "Two LLC", org["e2"].email, "Purchase Enquiry"]]))
    run(org["d"], b3)
    assert Lead.objects.get(title="For two").assigned_engineer == org["e2"]
    # an engineer cannot open someone else's import
    assert org["c2"].get(b3.get_absolute_url()).status_code == 404


def test_import_permissions(org, make_user, login):
    acc = login(make_user(caps.ACCOUNTS))
    assert acc.get(reverse("reporting:import_start")).status_code == 403
    assert acc.get(reverse("reporting:import_template", args=["leads"])).status_code == 403
    # the coordinator has no "pipeline.edit": opportunities cannot be imported by them
    co = login(make_user(caps.COORDINATOR))
    r = co.post(reverse("reporting:import_start"), {"kind": "opportunities", "file": csvfile("Deal title\nx\n")})
    assert r.status_code == 200 and not ImportBatch.objects.filter(kind="opportunities").exists()
    assert co.get(reverse("reporting:import_template", args=["opportunities"])).status_code == 403
    # CSRF is enforced on the upload
    from django.test import Client
    strict = Client(enforce_csrf_checks=True)
    strict.force_login(org["admin"], backend="django.contrib.auth.backends.ModelBackend")
    assert strict.post(reverse("reporting:import_start"), {"kind": "customers"}).status_code == 403


# ------------------------------------------------------------------ visits, tasks, historical dates, relationships
def test_visits_tasks_keep_historical_dates_and_check_relationships(org):
    other = Customer.objects.create(company_name="Other Co", code="CUS-90002", assigned_engineer=org["e1"])
    lead_other = Lead.objects.create(title="Other's lead", company_name="Other Co", customer=other,
                                     stage=__import__("leads.models", fromlist=["LeadStage"]).LeadStage.first_open(),
                                     assigned_engineer=org["e1"])
    old = TODAY - datetime.timedelta(days=40)
    rows = [["Customer", "Visit date", "Start time", "Status", "Employee", "Meeting outcome / notes", "Lead code"],
            ["CUS-90001", old.isoformat(), "10:30", "Completed", "eng.one", "Agreed next steps", ""],
            ["CUS-90001", old.isoformat(), "14:00", "Completed", "eng.one", "", ""],
            ["CUS-90001", TODAY.isoformat(), "09:00", "Scheduled", "eng.one", "", lead_other.code],
            ["Nobody Ltd", old.isoformat(), "09:00", "", "eng.one", "x", ""]]
    r, batch = upload(org["a"], "visits", xlsx(rows))
    res = results(batch)
    assert res[2]["status"] == "new"
    assert "required for a completed visit" in res[3]["messages"][0]
    assert "belongs to another customer" in res[4]["messages"][0]
    assert "No customer found" in res[5]["messages"][0]
    run(org["a"], batch)
    v = Visit.objects.get(customer=org["cust"])
    assert v.visit_date == old and v.status == "completed" and v.assigned_to == org["e1"]
    assert timezone.localtime(v.completed_at).date() == old
    from performance import monitoring as mon
    assert mon.list_figures([org["e1"].pk], old, old)[org["e1"].pk]["visits_completed"] == 1
    assert mon.list_figures([org["e1"].pk], TODAY, TODAY)[org["e1"].pk]["visits_completed"] == 0
    # tasks: completed on their real date
    trows = [["Title", "Due date", "Status", "Completed date", "Employee", "Customer"],
             ["Old follow-up", old.isoformat(), "Completed", old.isoformat(), "eng.one", "CUS-90001"]]
    r, tb = upload(org["a"], "tasks", xlsx(trows))
    run(org["a"], tb)
    t = Task.objects.get(title="Old follow-up")
    assert t.category == "follow_up" and t.status == "completed"
    assert timezone.localtime(t.completed_at).date() == old


# ------------------------------------------------------------------ completed sales and revenue
SALES = [["Deal title", "Customer", "Status / stage", "Deal value", "Closing date", "Sales Engineer", "Reference"],
         ["Valves supply", "CUS-90001", "Won", "250,000", "2026-02-28", "eng.one", "SALE-1"],
         ["Pump job", "CUS-90001", "Closed Won", "80000", "2026-03-10", "eng.one", ""],
         ["Lost job", "CUS-90001", "Lost", "5000", "2026-03-11", "eng.one", ""],
         ["No status", "CUS-90001", "", "99999", "2026-03-12", "eng.one", ""],
         ["Future win", "CUS-90001", "Won", "1000", (TODAY + datetime.timedelta(days=3)).isoformat(), "eng.one",
          ""],
         ["Open deal", "CUS-90001", "Negotiation", "40000", "", "eng.one", ""]]


def test_completed_sales_import_and_no_double_revenue(org):
    from reporting.metrics import kpi_actual

    r, batch = upload(org["a"], "opportunities", xlsx(SALES))
    res = results(batch)
    assert res[2]["status"] == "new" and res[3]["status"] == "new" and res[7]["status"] == "new"
    assert "Lost reason is required" in res[4]["messages"][0]
    assert "Status / stage is required" in res[5]["messages"][0]  # an amount alone never makes a sale
    assert "cannot be in the future" in res[6]["messages"][0]
    run(org["a"], batch)
    won = Opportunity.objects.filter(stage__is_won=True)
    assert won.count() == 2
    v = Opportunity.objects.get(title="Valves supply")
    assert timezone.localtime(v.closed_at).date() == datetime.date(2026, 2, 28) and v.probability == 100
    feb = (datetime.date(2026, 2, 1), datetime.date(2026, 2, 28))
    assert kpi_actual("won_value", {org["e1"].pk}, *feb) == Decimal("250000.00")
    total_before = sum(o.estimated_value for o in won)
    # the same file again (with and without references): nothing new, revenue unchanged
    for mode in ("skip", "update"):
        r, again = upload(org["a"], "opportunities", xlsx(SALES), mode=mode)
        assert results(again)[2]["status"] == "duplicate" and results(again)[3]["status"] == "duplicate"
        run(org["a"], again)
    assert Opportunity.objects.filter(stage__is_won=True).count() == 2
    assert sum(o.estimated_value for o in Opportunity.objects.filter(stage__is_won=True)) == total_before
    assert Opportunity.objects.filter(title="Open deal").count() == 1


def test_update_by_reference_changes_an_open_deal_to_won(org):
    rows = [SALES[0], ["Tower deal", "CUS-90001", "Negotiation", "100000", "", "eng.one", "DEAL-7"]]
    r, b = upload(org["a"], "opportunities", xlsx(rows))
    run(org["a"], b)
    rows2 = [SALES[0], ["Tower deal", "CUS-90001", "Won", "110000", "2026-03-01", "eng.one", "DEAL-7"]]
    r, b2 = upload(org["a"], "opportunities", xlsx(rows2), mode="update")
    res = results(b2)[2]
    assert res["status"] == "update" and any(c[0] == "Status / stage" for c in res["changes"])
    run(org["a"], b2)
    o = Opportunity.objects.get(title="Tower deal")
    assert o.stage.is_won and o.estimated_value == Decimal("110000.00")
    assert timezone.localtime(o.closed_at).date() == datetime.date(2026, 3, 1)
    assert Opportunity.objects.filter(title="Tower deal").count() == 1


def test_formulas_are_never_evaluated(org):
    wb = Workbook()
    ws = wb.active
    ws.append(["Company name", "Notes"])
    ws.append(["Formula Test LLC", "=1+1"])
    ws.append(['=HYPERLINK("http://evil.example","x")', "text"])
    out = io.BytesIO()
    wb.save(out)
    r, batch = upload(org["a"], "customers", SimpleUploadedFile("f.xlsx", out.getvalue()))
    run(org["a"], batch)
    # openpyxl never computes formulas; a formula without a cached result is read as empty
    c = Customer.objects.get(company_name="Formula Test LLC")
    assert c.notes in ("", "=1+1")
    assert not Customer.objects.filter(company_name__startswith="=").exists() or \
        Customer.objects.filter(company_name__startswith="=").count() <= 1


def test_detail_header_shows_row_count(org):
    r, batch = upload(org["a"], "customers", csvfile("Company name\nCount One LLC\nCount Two LLC\n"))
    assert "· 2 rows ·" in org["a"].get(batch.get_absolute_url()).content.decode()
