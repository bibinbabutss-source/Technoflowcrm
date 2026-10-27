"""Module-level checks: permissions per role, record scope, documents, imports, revisions, KPIs, MFA."""
import io
from decimal import Decimal

import pyotp
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone

from accounts import capabilities as caps
from core.money import LineInput, calculate, split_by_percentages
from customers.models import Customer
from documents.models import Document

pytestmark = pytest.mark.django_db

LIST_PAGES = {
    # url name: capability needed
    "customers:list": "customers.view", "leads:list": "leads.view", "leads:kanban": "leads.view",
    "visits:list": "visits.view", "visits:calendar": "visits.view", "visits:schedule": "visits.view",
    "sales:pipeline": "pipeline.view", "sales:kanban": "pipeline.view", "sales:forecast": "pipeline.view",
    "sales:projects": "projects.view", "quotations:list": "quotations.view", "orders:list": "orders.view",
    "invoicing:list": "invoices.view", "invoicing:payments": "collections.view",
    "invoicing:collections": "collections.view", "performance:overview": "targets.view",
    "tasks:list": "tasks.view", "tasks:calendar": "tasks.view", "communications:list": "communications.view",
    "documents:list": "documents.view", "reporting:index": "reports.view", "notifications:list": "notifications.view",
    "core:settings_home": "settings.manage", "core:settings_workflow": "settings.manage",
    "core:tax_rates": "settings.manage", "core:audit_log": "audit.view", "accounts:user_list": "users.manage",
    "reporting:import_start": "data.import", "performance:kpis": "targets.manage",
    "customers:create": "customers.create", "leads:create": "leads.create", "quotations:create": "quotations.edit",
    "invoicing:payment_create": "payments.manage",
}


@pytest.mark.parametrize("role", [caps.ADMIN, caps.DIRECTOR, caps.ACCOUNTS, caps.ENGINEER, caps.COORDINATOR])
def test_every_page_allows_or_refuses_by_capability(role, make_user, login):
    user = make_user(role)
    client = login(user)
    codes = user.profile.capability_codes()
    problems = []
    for name, cap in LIST_PAGES.items():
        url = reverse(name)
        r = client.get(url, follow=False)
        if r.status_code == 302 and name == "tasks:list":
            r = client.get(r["Location"])
        expected = 200 if cap in codes else 403
        if r.status_code != expected:
            problems.append(f"{name}: got {r.status_code}, expected {expected}")
    assert not problems, problems


def test_engineer_cannot_open_another_engineers_customer(make_user, login):
    a, b = make_user(caps.ENGINEER), make_user(caps.ENGINEER)
    cust = Customer.objects.create(company_name="Scope Test LLC", assigned_engineer=a)
    assert login(a).get(cust.get_absolute_url()).status_code == 200
    assert login(b).get(cust.get_absolute_url()).status_code == 404
    assert "Scope Test LLC" not in login(b).get(reverse("customers:list")).content.decode()


def test_server_rejects_actions_without_capability(make_user, login):
    eng = make_user(caps.ENGINEER)
    cust = Customer.objects.create(company_name="Delete Test LLC", assigned_engineer=eng)
    client = login(eng)
    r = client.post(reverse("customers:delete", args=[cust.pk]))
    assert r.status_code == 403
    assert Customer.objects.filter(pk=cust.pk).exists()
    assert client.post(reverse("core:run_reminders")).status_code == 403


def test_duplicate_trn_is_blocked(make_user, login):
    admin = make_user(caps.ADMIN)
    Customer.objects.create(company_name="Original LLC", trn="100200300400500")
    r = login(admin).post(reverse("customers:create"), {"company_name": "Another Name LLC",
                                                         "trn": "100200300400500", "status": "active",
                                                         "country": "United Arab Emirates"})
    assert r.status_code == 200
    assert Customer.objects.filter(trn="100200300400500").count() == 1


# --------------------------------------------------------------- documents
def _upload(client, cust, name, content, ctype="application/pdf"):
    return client.post(reverse("documents:upload"), {
        "target_type": "customers.customer", "target_id": cust.pk, "title": name,
        "file": SimpleUploadedFile(name, content, content_type=ctype)})


def test_document_upload_validation_and_private_download(make_user, login, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    eng, other = make_user(caps.ENGINEER), make_user(caps.ENGINEER)
    cust = Customer.objects.create(company_name="Docs LLC", assigned_engineer=eng)
    c = login(eng)
    _upload(c, cust, "drawing.pdf", b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF")
    doc = Document.objects.get(title="drawing.pdf")
    # executable / script types are refused, and a fake PDF (wrong magic bytes) too
    _upload(c, cust, "evil.exe", b"MZ\x90\x00", "application/octet-stream")
    _upload(c, cust, "evil.html", b"<script>alert(1)</script>", "text/html")
    _upload(c, cust, "fake.pdf", b"MZ\x90\x00 not a pdf")
    assert Document.objects.count() == 1
    # stored under a random name, outside static/public folders
    assert doc.file.name != "drawing.pdf" and "drawing" not in doc.file.name
    r = c.get(reverse("documents:download", args=[doc.pk]))
    assert r.status_code == 200
    assert r["Content-Disposition"].startswith("attachment")
    assert r["X-Content-Type-Options"] == "nosniff"
    assert login(other).get(reverse("documents:download", args=[doc.pk])).status_code == 404


# ------------------------------------------------------------------ import
def test_import_previews_and_never_overwrites(make_user, login):
    admin = make_user(caps.ADMIN)
    Customer.objects.create(company_name="Existing Co LLC", email="x@example.com")
    c = login(admin)
    csv = ("Company,Email,City\nExisting Co LLC,x@example.com,Dubai\nNew Import LLC,new@example.com,Sharjah\n"
           "New Import LLC,new@example.com,Sharjah\n,missing@example.com,Dubai\n").encode()
    r = c.post(reverse("reporting:import_start"), {"kind": "customers",
                                                   "file": SimpleUploadedFile("c.csv", csv, "text/csv")})
    assert r.status_code == 302
    detail = r["Location"]
    from reporting.models import ImportBatch

    batch = ImportBatch.objects.get()
    assert batch.mapping["company_name"] == 0 and batch.mapping["email"] == 1
    r = c.post(detail, {"map_company_name": "0", "map_email": "1", "map_city": "2"})
    batch.refresh_from_db()
    assert batch.summary == {"total": 4, "ok": 1, "duplicates": 2, "errors": 1}
    c.post(reverse("reporting:import_run", args=[batch.pk]))
    assert Customer.objects.filter(company_name="New Import LLC").count() == 1
    assert Customer.objects.get(company_name="Existing Co LLC").city == ""  # not overwritten
    r = c.get(reverse("reporting:import_errors", args=[batch.pk]))
    assert r.status_code == 200 and b"Company name is required" in r.content


# ------------------------------------------------------------ calculations
def test_money_calculation_header_discount_and_mixed_vat():
    t = calculate([LineInput(Decimal(2), Decimal("100.00"), Decimal(10), Decimal(5)),
                   LineInput(Decimal(1), Decimal("50.00"), Decimal(0), Decimal(0))], Decimal("23.00"))
    assert t.subtotal == Decimal("230.00")
    assert t.taxable == Decimal("207.00")
    assert t.tax == Decimal("8.10")  # 5% of (180 − 18 header-discount share)
    assert t.grand_total == Decimal("215.10")
    parts = split_by_percentages(Decimal("100.00"), [Decimal(30), Decimal(40), Decimal(30)])
    assert sum(parts) == Decimal("100.00")


def test_revision_keeps_previous_version_unchanged(make_user, login):
    from quotations import services
    from quotations.models import QuotationLine
    from core.models import TaxRate

    eng, director = make_user(caps.ENGINEER), make_user(caps.DIRECTOR)
    cust = Customer.objects.create(company_name="Rev LLC", assigned_engineer=eng)
    q, rev = services.create_quotation(customer=cust, opportunity=None, salesperson=eng, currency="AED", user=eng)
    QuotationLine.objects.create(revision=rev, description="Item", quantity=1, unit_price=1000,
                                 tax_rate=TaxRate.default())
    services.recalculate(rev)
    services.submit(rev, eng)
    rev.refresh_from_db()
    services.approve(rev, director)
    login(eng).post(reverse("quotations:revise", args=[q.pk]), {"reason": "Customer asked for more"})
    q.refresh_from_db()
    assert q.current_version == 2
    v1 = q.revisions.get(version=1)
    assert v1.status == "superseded" and v1.is_locked and v1.grand_total == Decimal("1050.00")
    assert q.revisions.get(version=2).lines.count() == 1


# --------------------------------------------------------------------- MFA
def test_two_step_sign_in(make_user, client):
    from tests.conftest import PASSWORD

    user = make_user(caps.ENGINEER)
    secret = pyotp.random_base32()
    user.profile.mfa_enabled, user.profile.mfa_secret = True, secret
    user.profile.save()
    r = client.post(reverse("accounts:login"), {"username": user.username, "password": PASSWORD})
    assert r.status_code == 302 and r["Location"] == reverse("accounts:mfa_verify")
    assert client.get(reverse("core:dashboard")).status_code == 302  # not signed in yet
    r = client.post(reverse("accounts:mfa_verify"), {"code": "000000"})
    assert r.status_code == 200
    r = client.post(reverse("accounts:mfa_verify"), {"code": pyotp.TOTP(secret).now()})
    assert r.status_code == 302
    assert client.get(reverse("core:dashboard")).status_code == 200


# --------------------------------------------------------------- KPIs
def test_targets_and_scorecard(make_user, login):
    from performance.models import KPIDefinition, SalesTarget
    from performance.services import scorecard

    director, eng = make_user(caps.DIRECTOR), make_user(caps.ENGINEER)
    kpi = KPIDefinition.objects.get(code="new_customers")
    month = timezone.localdate().replace(day=1)
    SalesTarget.objects.create(employee=eng, kpi=kpi, period_type="month", period_start=month, target_value=4)
    Customer.objects.create(company_name="KPI One LLC", assigned_engineer=eng)
    card = scorecard(eng, "month", timezone.localdate(), compare=False)
    row = next(r for r in card["rows"] if r["kpi"].code == "new_customers")
    assert row["actual"] == 1 and row["pct"] == Decimal("25.0")
    assert login(director).get(reverse("performance:overview") + f"?employee={eng.pk}").status_code == 200
    assert login(eng).get(reverse("performance:set_targets")).status_code == 403


# ------------------------------------------------------------------ exports
def test_exports_neutralise_spreadsheet_formulas_and_need_permission(make_user, login):
    from openpyxl import load_workbook

    admin, eng = make_user(caps.ADMIN), make_user(caps.ENGINEER)
    Customer.objects.create(company_name='=HYPERLINK("http://evil.example","x")', assigned_engineer=eng)
    c = login(admin)
    r = c.get(reverse("customers:list") + "?export=csv")
    assert r.status_code == 200
    assert "'=HYPERLINK" in r.content.decode("utf-8-sig")
    r = c.get(reverse("customers:list") + "?export=xlsx")
    ws = load_workbook(io.BytesIO(r.content)).active
    values = [cell.value for row in ws.iter_rows() for cell in row if cell.value]
    assert any(str(v).startswith("'=HYPERLINK") for v in values)
    assert not any(str(v).startswith("=") for v in values)
    if "data.export" not in eng.profile.capability_codes():
        assert login(eng).get(reverse("customers:list") + "?export=csv").status_code == 403
    r = c.get(reverse("reporting:report", args=["customer_list"]) + "?export=xlsx")
    assert r.status_code == 200
    values = [cell.value for row in load_workbook(io.BytesIO(r.content)).active.iter_rows() for cell in row]
    assert not any(isinstance(v, str) and v.startswith("=") for v in values)


@pytest.mark.parametrize("role", [caps.DIRECTOR, caps.ACCOUNTS, caps.ENGINEER])
def test_every_report_renders_and_exports(role, make_user, login):
    from reporting.reports import REPORTS

    user = make_user(role)
    codes = user.profile.capability_codes()
    c = login(user)
    checked = 0
    for rep in REPORTS:
        url = reverse("reporting:report", args=[rep.key])
        r = c.get(url)
        if rep.capability not in codes:
            assert r.status_code == 403, rep.key
            continue
        assert r.status_code == 200, rep.key
        for fmt in ("csv", "xlsx", "pdf"):
            r = c.get(url + f"?export={fmt}")
            expected = 200 if (fmt == "pdf" or "data.export" in codes) else 403
            assert r.status_code == expected, (rep.key, fmt, r.status_code)
        checked += 1
    assert checked > 0
