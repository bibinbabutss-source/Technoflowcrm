"""
Controlled import of customers and leads from Excel (.xlsx) or CSV.

1. Upload — file type and size checked, at most 2,000 rows.
2. Map columns — the CRM suggests a match for each field from the column headings.
3. Preview — every row is validated; duplicates (existing records or repeated rows in the file) are flagged.
4. Import — only valid, non-duplicate rows are created. Existing records are NEVER changed or overwritten.
5. Summary — counts and a downloadable error report.
"""
import csv
import io
import re
from decimal import Decimal, InvalidOperation

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction

from accounts.permissions import has_capability
from core.activity import record
from core.audit import log_action
from core.models import LookupValue

MAX_ROWS = 2000
MAX_BYTES = 5 * 1024 * 1024

# field, label, required, synonyms
FIELDS = {
    "customers": [
        ("company_name", "Company name", True, ["company", "company name", "customer", "client", "client name",
                                                "name", "organisation", "organization"]),
        ("category", "Category", False, ["category", "client type", "customer type", "type"]),
        ("industry", "Industry", False, ["industry", "sector"]),
        ("status", "Status", False, ["status"]),
        ("contact_person", "Contact person", False, ["contact", "contact person", "contact name"]),
        ("designation", "Designation", False, ["designation", "title", "position", "job title"]),
        ("phone", "Phone", False, ["phone", "mobile", "telephone", "tel", "primary phone"]),
        ("alt_phone", "Alternative phone", False, ["alternative phone", "alt phone", "phone 2", "other phone"]),
        ("email", "Email", False, ["email", "e-mail", "email address"]),
        ("website", "Website", False, ["website", "web", "url"]),
        ("billing_address", "Billing address", False, ["billing address", "address", "invoice address"]),
        ("office_address", "Office address", False, ["office address", "office"]),
        ("city", "Emirate / city", False, ["city", "emirate", "location"]),
        ("country", "Country", False, ["country"]),
        ("trn", "TRN", False, ["trn", "tax registration number", "vat number", "tax number"]),
        ("source", "Source", False, ["source", "lead source"]),
        ("assigned_engineer", "Sales Engineer (username or email)", False, ["sales engineer", "assigned to",
                                                                            "owner", "salesperson"]),
        ("notes", "Notes", False, ["notes", "remarks", "comments"]),
    ],
    "leads": [
        ("title", "Lead title", True, ["title", "lead title", "lead", "project", "subject", "enquiry"]),
        ("company_name", "Company name", True, ["company", "company name", "customer", "client", "client name"]),
        ("contact_person", "Contact person", False, ["contact", "contact person", "contact name"]),
        ("designation", "Designation", False, ["designation", "position", "job title"]),
        ("email", "Email", False, ["email", "e-mail"]),
        ("phone", "Phone", False, ["phone", "mobile", "telephone", "tel"]),
        ("city", "Emirate / city", False, ["city", "emirate", "location"]),
        ("source", "Source", False, ["source", "lead source"]),
        ("interest", "Service / product interest", False, ["interest", "service", "product", "requirement"]),
        ("estimated_value", "Estimated value", False, ["value", "estimated value", "potential value", "amount",
                                                       "potential value (aed)", "estimated value (aed)"]),
        ("priority", "Priority (Low / Medium / High)", False, ["priority"]),
        ("assigned_engineer", "Sales Engineer (username or email)", False, ["sales engineer", "assigned to",
                                                                            "owner", "assigned"]),
        ("qualification_notes", "Notes", False, ["notes", "remarks", "comments", "discussion summary"]),
    ],
}


def norm(text):
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower()).strip()


def parse_file(f):
    name = f.name.lower()
    if f.size > MAX_BYTES:
        raise ValidationError("The file is larger than 5 MB.")
    if name.endswith(".csv"):
        raw = f.read()
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = list(csv.reader(io.StringIO(text), dialect))
    elif name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(f, read_only=True, data_only=True)
        except Exception:
            raise ValidationError("This is not a valid Excel .xlsx file.")
        ws = wb.worksheets[0]
        reader = [[("" if c is None else str(c).strip()) for c in row] for row in ws.iter_rows(values_only=True)]
    else:
        raise ValidationError("Upload a .xlsx or .csv file.")
    reader = [r for r in reader if any(str(c).strip() for c in r)]
    if not reader:
        raise ValidationError("The file has no rows.")
    headers = [str(h).strip() or f"Column {i + 1}" for i, h in enumerate(reader[0])]
    rows = []
    for r in reader[1:]:
        cells = [str(c).strip() if c is not None else "" for c in r]
        cells += [""] * (len(headers) - len(cells))
        rows.append(cells[:len(headers)])
    if not rows:
        raise ValidationError("The file has a heading row but no data rows.")
    if len(rows) > MAX_ROWS:
        raise ValidationError(f"The file has {len(rows)} rows; import at most {MAX_ROWS} at a time.")
    return headers, rows


def guess_mapping(kind, headers):
    mapping = {}
    normalized = {norm(h): i for i, h in enumerate(headers)}
    for fname, _label, _req, synonyms in FIELDS[kind]:
        for syn in [fname.replace("_", " ")] + synonyms:
            if norm(syn) in normalized:
                mapping[fname] = normalized[norm(syn)]
                break
    return mapping


def _lookup(category, value):
    if not value:
        return None, None
    lv = LookupValue.objects.filter(category=category, value__iexact=value.strip()).first()
    if lv is None:
        return None, f"'{value}' is not in the {category.replace('_', ' ')} list (Settings > Drop-down lists)"
    return lv, None


def _user(value):
    if not value:
        return None, None
    User = get_user_model()
    u = User.objects.filter(is_active=True, profile__isnull=False).filter(username__iexact=value.strip()).first() or \
        User.objects.filter(is_active=True, profile__isnull=False, email__iexact=value.strip()).first()
    return (u, None) if u else (None, f"No active employee with username or email '{value}'")


def build_rows(batch, user):
    """Validate every row with the current mapping. Returns list of dicts: data, errors, duplicate."""
    from customers.models import normalize_name
    from customers.services import find_duplicates
    from leads.services import find_lead_duplicates

    kind = batch.kind
    fields = FIELDS[kind]
    can_assign = has_capability(user, "leads.assign" if kind == "leads" else "records.assign")
    seen = set()
    out = []
    for idx, raw in enumerate(batch.rows, start=2):
        data, errors, warnings = {}, [], []
        for fname, label, required, _ in fields:
            col = batch.mapping.get(fname)
            value = raw[col].strip() if col is not None and col < len(raw) else ""
            if required and not value:
                errors.append(f"{label} is required")
            data[fname] = value
        if data.get("email"):
            try:
                validate_email(data["email"])
            except ValidationError:
                errors.append(f"'{data['email']}' is not a valid email")
        for fname, category in (("category", "customer_category"), ("industry", "industry"),
                                ("source", "lead_source")):
            if fname in data and data[fname]:
                obj, err = _lookup(category, data[fname])
                if err:
                    errors.append(err)
                data[fname] = obj
            elif fname in data:
                data[fname] = None
        if "assigned_engineer" in data:
            u, err = _user(data["assigned_engineer"])
            if err:
                errors.append(err)
            if u and not can_assign and u != user:
                warnings.append("You cannot assign to others; you will be the owner")
                u = user
            data["assigned_engineer"] = u or (user if not can_assign else None)
        if kind == "customers" and data.get("status"):
            mapping = {"active": "active", "prospect": "prospect", "lead": "prospect", "inactive": "inactive",
                       "on hold": "on_hold"}
            st = mapping.get(norm(data["status"]))
            if st is None:
                errors.append(f"Unknown status '{data['status']}' (use Active, Prospect, Inactive or On hold)")
            data["status"] = st or "prospect"
        elif kind == "customers":
            data["status"] = "prospect"
        if kind == "leads":
            if data.get("estimated_value"):
                try:
                    v = Decimal(re.sub(r"[^\d.\-]", "", data["estimated_value"]) or "0")
                    if v < 0:
                        raise InvalidOperation
                    data["estimated_value"] = v
                except InvalidOperation:
                    errors.append(f"'{data['estimated_value']}' is not a valid amount")
                    data["estimated_value"] = Decimal(0)
            else:
                data["estimated_value"] = Decimal(0)
            pr = norm(data.get("priority"))
            data["priority"] = pr if pr in ("low", "medium", "high") else "medium"
        duplicate = ""
        if not errors:
            if kind == "customers":
                key = normalize_name(data["company_name"])
                dups = find_duplicates(company_name=data["company_name"], email=data.get("email"),
                                       phone=data.get("phone"), trn=data.get("trn"))
                if dups:
                    duplicate = f"Already exists: {dups[0]['customer']} ({', '.join(dups[0]['reasons'])})"
            else:
                key = (normalize_name(data["company_name"]), norm(data["title"]))
                dups = find_lead_duplicates(company_name=data["company_name"], email=data.get("email"),
                                            phone=data.get("phone"))
                if dups:
                    duplicate = f"Similar open lead: {dups[0].code} {dups[0].title}"
            if not duplicate and key in seen:
                duplicate = "Repeated in this file"
            seen.add(key)
        out.append({"row": idx, "data": data, "errors": errors, "warnings": warnings, "duplicate": duplicate,
                    "raw": raw})
    return out


def preview(batch, user):
    rows = build_rows(batch, user)
    batch.results = [{"row": r["row"], "status": "error" if r["errors"] else "duplicate" if r["duplicate"] else "ok",
                      "messages": r["errors"] + ([r["duplicate"]] if r["duplicate"] else []) + r["warnings"]}
                     for r in rows]
    batch.summary = {"total": len(rows), "ok": sum(1 for r in batch.results if r["status"] == "ok"),
                     "duplicates": sum(1 for r in batch.results if r["status"] == "duplicate"),
                     "errors": sum(1 for r in batch.results if r["status"] == "error")}
    batch.status = batch.Status.PREVIEW
    batch.save()
    return rows


@transaction.atomic
def run_import(batch, user):
    from customers.models import Customer
    from leads.models import Lead, LeadStage
    from leads.services import record_assignment

    if batch.status != batch.Status.PREVIEW:
        raise ValidationError("Preview the import before running it.")
    rows = build_rows(batch, user)  # validate again: data may have changed since the preview
    created = 0
    for r in rows:
        if r["errors"] or r["duplicate"]:
            continue
        d = r["data"]
        if batch.kind == "customers":
            obj = Customer(**{k: v for k, v in d.items() if v not in (None, "")})
            obj.save()
            record(obj, "created", f"Imported from {batch.file_name}")
        else:
            obj = Lead(stage=LeadStage.first_open(), **{k: v for k, v in d.items() if v not in (None, "")})
            obj.save()
            record(obj, "created", f"Imported from {batch.file_name}")
            record_assignment(obj, note="Imported")
        created += 1
    preview(batch, user)
    batch.summary["created"] = created
    batch.status = batch.Status.IMPORTED
    batch.save()
    log_action("other", batch, message=f"Imported {created} {batch.kind} from {batch.file_name} "
                                      f"({batch.summary['duplicates']} duplicates and {batch.summary['errors']} "
                                      f"errors skipped)")
    return created
