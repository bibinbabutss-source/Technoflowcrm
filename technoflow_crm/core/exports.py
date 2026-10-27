"""CSV and Excel export helpers shared by every list and report."""
import csv
import datetime
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe_text(value):
    """Prevent spreadsheet formula injection (a cell starting with = runs as a formula in Excel)."""
    if value is None:
        return ""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def _plain(value):
    if isinstance(value, datetime.datetime):
        return timezone.localtime(value).strftime("%d-%b-%Y %H:%M") if timezone.is_aware(value) else value
    if isinstance(value, datetime.date):
        return value.strftime("%d-%b-%Y")
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return value


def csv_response(filename, headers, rows):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
    response.write("﻿")  # BOM so Excel opens UTF-8 correctly
    writer = csv.writer(response)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([safe_text(_plain(v)) if not isinstance(v, (int, float, Decimal)) else v for v in row])
    return response


def xlsx_response(filename, headers, rows, *, title=None, subtitle=None, totals=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = (title or filename)[:31]
    r = 1
    if title:
        ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=14)
        r += 1
    if subtitle:
        ws.cell(row=r, column=1, value=subtitle).font = Font(italic=True, color="5F6062")
        r += 1
    if title or subtitle:
        r += 1
    header_row = r
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = Font(bold=True, color="23262A")
        cell.fill = PatternFill("solid", fgColor="E9EDB0")
        cell.alignment = Alignment(vertical="center")
    widths = [len(str(h)) for h in headers]
    for row in rows:
        r += 1
        for c, value in enumerate(row, 1):
            value = _plain(value)
            if isinstance(value, Decimal):
                value = float(value)
            elif isinstance(value, str):
                value = safe_text(value)
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, float):
                cell.number_format = "#,##0.00"
            widths[c - 1] = max(widths[c - 1], min(len(str(value)) if value is not None else 0, 60))
    if totals:
        r += 1
        for c, value in enumerate(totals, 1):
            if value in (None, ""):
                continue
            cell = ws.cell(row=r, column=c, value=float(value) if isinstance(value, Decimal) else value)
            cell.font = Font(bold=True)
            if isinstance(value, (Decimal, float)):
                cell.number_format = "#,##0.00"
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, w + 2)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response
