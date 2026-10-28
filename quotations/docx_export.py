"""
Editable Word (.docx) version of the quotation, in the same order and wording as the PDF:
cover letter, price schedule (column titles repeat on every page), commercial terms and conditions.
"""
import re
from io import BytesIO

from django.contrib.staticfiles import finders

from . import document


def _bold_runs(paragraph, text, size=None):
    """Add text to a paragraph, turning **bold** markers into bold runs."""
    from docx.shared import Pt

    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text or "")):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = i % 2 == 1
        if size:
            run.font.size = Pt(size)
    return paragraph


def _paragraphs(cell_or_doc, text, size=10, first=None):
    blocks = [b for b in re.split(r"\n\s*\n", text or "") if b.strip()]
    out = []
    for n, block in enumerate(blocks):
        p = first if (n == 0 and first is not None) else cell_or_doc.add_paragraph()
        for k, line in enumerate(block.split("\n")):
            if k:
                p.add_run().add_break()
            _bold_runs(p, line, size)
        out.append(p)
    return out


def _set_cell_border(cell, **kwargs):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            el = OxmlElement(f"w:{edge}")
            for key, val in kwargs[edge].items():
                el.set(qn(f"w:{key}"), str(val))
            borders.append(el)


def _repeat_header(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def build_docx(q, rev):
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Cm, Pt

    d = document.build(q, rev)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.3)
    sec.top_margin, sec.bottom_margin = Cm(1.0), Cm(1.2)

    # letterhead (header) and footer, on every page
    logo = finders.find("img/technoflow-logo-full.png")
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if logo:
        hp.add_run().add_picture(logo, width=Cm(3.9))
    for line in (d["tagline"], d["arabic_name"], d["legal_name"], f"TRN:{d['trn']}" if d["trn"] else ""):
        if line:
            p = sec.header.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(line)
            r.font.size = Pt(6.5)
    fp = sec.footer.paragraphs[0]
    fp.add_run(d["footer_tagline"]).font.size = Pt(8)
    p = sec.footer.add_paragraph()
    r = p.add_run(d["footer_website"])
    r.bold, r.font.size = True, Pt(8)
    p = sec.footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{d['footer_business_lines']}\n{d['footer_contact']}")
    r.font.size = Pt(7)

    # 1. cover letter
    if d["is_draft"]:
        p = doc.add_paragraph()
        r = p.add_run("DRAFT — NOT APPROVED")
        r.bold = True
    doc.add_paragraph(d["date"].strftime("%d/%m/%Y"))
    p = doc.add_paragraph()
    p.add_run(d["customer_name"]).bold = True
    if d["customer_address"]:
        doc.add_paragraph(d["customer_address"])
    info = doc.add_table(rows=0, cols=2)
    rows = [("K/Attn:", d["attention"], False), ("Subject :", d["subject"], True), ("Quote. Ref :", d["number"], False),
            ("Project:", d["project"], False), ("Revision Number", str(d["revision_number"]), False)]
    for label, value, bold in rows:
        if not value and label in ("K/Attn:", "Project:"):
            continue
        cells = info.add_row().cells
        cells[0].width, cells[1].width = Cm(5.2), Cm(12.8)
        cells[0].paragraphs[0].add_run(label).bold = bold
        cells[1].paragraphs[0].add_run(value or "").bold = bold
    doc.add_paragraph()
    _paragraphs(doc, d["letter_body"], 10)
    doc.add_paragraph()
    p = doc.add_paragraph()
    for k, line in enumerate((d["letter_closing"] or "").split("\n")):
        if k:
            p.add_run().add_break()
        p.add_run(line).font.size = Pt(11)
    for _ in range(3):
        doc.add_paragraph()
    sig = doc.add_table(rows=1, cols=2)
    left = sig.rows[0].cells[0].paragraphs[0]
    if d["signatory_name"]:
        left.add_run(d["signatory_name"]).add_break()
    left.add_run(". . . . . . . . . . . . . . . . . . . . . . . . . . .").add_break()
    left.add_run(d["signatory_designation"]).font.size = Pt(11)
    right = sig.rows[0].cells[1].paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.add_run(d["signature_label"]).font.size = Pt(11)

    # 2. price schedule
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PRICE SCHEDULE")
    r.bold, r.underline, r.font.size = True, True, Pt(14)
    head = doc.add_table(rows=3, cols=4)
    pairs = [("Customer", d["customer_name"], "Date", d["date"].strftime("%d/%m/%Y")),
             ("Project", d["project"], "Currency", d["currency_name"]),
             ("Consultant", d["consultant"], "Quote No.", d["number"])]
    for row, (k1, v1, k2, v2) in zip(head.rows, pairs):
        for cell, text, bold in zip(row.cells, (k1, f": {v1}", k2, f": {v2}"), (True, False, True, False)):
            run = cell.paragraphs[0].add_run(text)
            run.bold, run.font.size = bold, Pt(8.5)
    if d["manufacturer"]:
        p = doc.add_paragraph()
        r = p.add_run(f"MANUFACTURER : {d['manufacturer']}")
        r.bold, r.font.size = True, Pt(8.5)
    titles = ["Sl NO.", "PART NO.", "ITEM DESCRIPTION", "QTY", "UNIT PRICE"] + (
        ["DISC."] if d["show_discount"] else []) + ["TOTAL PRICE"]
    widths = [1.1, 2.8, 7.2, 1.4, 2.4] + ([1.2] if d["show_discount"] else []) + [2.5]
    table = doc.add_table(rows=1, cols=len(titles))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, title, w in zip(table.rows[0].cells, titles, widths):
        cell.width = Cm(w)
        run = cell.paragraphs[0].add_run(title)
        run.bold, run.font.size = True, Pt(8.5)
        _set_cell_border(cell, bottom={"val": "single", "sz": 6, "color": "000000"})
    _repeat_header(table.rows[0])
    for n, ln in enumerate(d["lines"], start=1):
        values = [str(n), ln.part_no, ln.description, f"{ln.quantity.normalize():f}" + (
            f" {ln.unit}" if d["show_unit"] and ln.unit else ""), f"{ln.unit_price:,.2f}"]
        if d["show_discount"]:
            values.append(f"{ln.discount_pct.normalize():f}%" if ln.discount_pct else "")
        values.append(f"{ln.line_net:,.2f}")
        cells = table.add_row().cells
        for i, (cell, value, w) in enumerate(zip(cells, values, widths)):
            cell.width = Cm(w)
            para = cell.paragraphs[0]
            para.add_run(value or "").font.size = Pt(8.5)
            if i == 3 or i == 0:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif i >= 4:
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    totals = []
    if not d["simple_total"]:
        totals.append(("SUB TOTAL", f"{d['subtotal']:,.2f}"))
        if d["header_discount"]:
            totals += [("DISCOUNT", f"-{d['header_discount']:,.2f}"), ("TOTAL AFTER DISCOUNT", f"{d['taxable']:,.2f}")]
        totals += [(f"VAT {rate.normalize():f}%", f"{tax:,.2f}") for rate, tax in d["tax_rows"]]
    totals.append(("TOTAL AMOUNT", f"{d['grand_total']:,.2f}"))
    tt = doc.add_table(rows=0, cols=2)
    tt.alignment = WD_TABLE_ALIGNMENT.RIGHT
    for label, value in totals:
        cells = tt.add_row().cells
        cells[0].width, cells[1].width = Cm(5), Cm(3.5)
        for cell, text in zip(cells, (label, value)):
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = para.add_run(text)
            run.bold = True
    p = doc.add_paragraph()
    r = p.add_run(f"Total Price: {d['currency_name']}  ")
    r.bold, r.font.size = True, Pt(8.5)
    p.add_run(d["amount_words"]).font.size = Pt(8.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(d["schedule_note"]).bold = True
    sigs = doc.add_table(rows=2, cols=3)
    for cell, name, role in zip(sigs.columns, (d["sales_engineer"], d["sales_in_charge"], d["approved_by"]),
                                ("Sales Engineer", "Sales In Charge", "Approved By")):
        cell.cells[0].paragraphs[0].add_run(name or "")
        cell.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.cells[1].paragraphs[0].add_run(role).bold = True
        cell.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 3. commercial terms
    if d["terms"]:
        new = doc.add_section(WD_SECTION.NEW_PAGE)
        new.left_margin = new.right_margin = Cm(2.2)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("COMMERCIAL TERMS AND CONDITIONS")
        r.bold, r.underline, r.font.size = True, True, Pt(13)
        terms = doc.add_table(rows=0, cols=2)
        terms.style = "Table Grid"
        for heading, content in d["terms"]:
            cells = terms.add_row().cells
            cells[0].width, cells[1].width = Cm(4.2), Cm(12.4)
            hp = cells[0].paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            hp.add_run(heading).bold = True
            _paragraphs(cells[1], content, 10, first=cells[1].paragraphs[0])

    out = BytesIO()
    doc.save(out)
    return out.getvalue()
