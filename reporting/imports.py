"""
Controlled import from Excel (.xlsx) or CSV: customers, sales leads, customer visits, follow-ups / tasks,
opportunities and completed sales.

1. Upload — the file's CONTENT is checked (a real .xlsx workbook or plain-text CSV), ≤ 5 MB, ≤ 2,000 rows,
   ≤ 60 columns. Formulas are never evaluated: the stored cell value is read as text.
2. Map columns — matched automatically from the headings, adjustable.
3. Preview — every row is validated (required fields, dates, numbers, emails, phones, choices, employees,
   customer / lead / opportunity references) and classified: NEW, UPDATE (reliable match, with the changes),
   DUPLICATE (skipped) or ERROR (skipped, with the Excel row number and the reason). Nothing is saved yet.
4. Import — only after confirmation, in ONE database transaction: if anything fails, nothing is saved.
   Existing records change only in an update mode, only on a reliable match, and only the fields given in
   the file (an empty cell never blanks a value).
5. Report — counts, and a downloadable CSV of every row that was not imported, with its reason.

Ownership: the employee named in the file when the importer may assign to them; empty → the importer, but
only when the importer is a Sales Engineer or Coordinator (never silently the Administrator).
Historical dates in the file are kept (lead date received, visit date, task completion, deal closing date).
"""
import csv
import datetime
import io
import logging

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone

from core.activity import record
from core.audit import log_action

from .import_kinds import KINDS, Context, fingerprint, norm, read_cell, ref_key

log = logging.getLogger(__name__)

MAX_ROWS = 2000
MAX_COLS = 60
MAX_BYTES = 5 * 1024 * 1024
OWNER_FIELD = {"customers": "assigned_engineer", "leads": "assigned_engineer", "visits": "assigned_to",
               "tasks": "assigned_to", "opportunities": "assigned_engineer"}
# kept for older code that imported FIELDS: (name, label, required, synonyms)
FIELDS = {k: [(f.name, f.label, f.required, f.synonyms) for f in v["fields"]] for k, v in KINDS.items()}


# ------------------------------------------------------------------ 1. reading the file
def parse_file(f):
    name = (f.name or "").lower()
    if f.size > MAX_BYTES:
        raise ValidationError("The file is larger than 5 MB.")
    head = f.read(8)
    f.seek(0)
    if name.endswith(".xlsx"):
        if head[:4] != b"PK\x03\x04":
            raise ValidationError("This is not a real Excel .xlsx file (its content does not match). "
                                  "Save it from Excel as 'Excel Workbook (.xlsx)'.")
        from openpyxl import load_workbook

        try:
            wb = load_workbook(f, read_only=True, data_only=True)
            ws = wb.worksheets[0]
            reader = []
            for row in ws.iter_rows(values_only=True):
                reader.append([_cell(c) for c in row[:MAX_COLS + 1]])
                if len(reader) > MAX_ROWS + 50:
                    break
            wb.close()
        except Exception:  # corrupt or unusual workbook: never show internal details
            raise ValidationError("The Excel file could not be read. Open it in Excel, save it again as .xlsx "
                                  "and retry.")
    elif name.endswith(".csv"):
        raw = f.read()
        if b"\x00" in raw[:4096] or head[:4] == b"PK\x03\x04":
            raise ValidationError("This is not a plain-text CSV file.")
        text = None
        for enc in ("utf-8-sig", "cp1252"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValidationError("The CSV file is not in a readable text encoding (use UTF-8).")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = []
        for row in csv.reader(io.StringIO(text), dialect):
            reader.append([str(c).strip() for c in row[:MAX_COLS + 1]])
            if len(reader) > MAX_ROWS + 50:
                break
    else:
        raise ValidationError("Upload an Excel .xlsx or a .csv file. Other file types are not accepted.")
    reader = [r for r in reader if any(str(c).strip() for c in r)]
    if not reader:
        raise ValidationError("The file has no rows.")
    if len(reader[0]) > MAX_COLS:
        raise ValidationError(f"The file has more than {MAX_COLS} columns. Remove the columns you don't need.")
    headers = [str(h).strip()[:100] or f"Column {i + 1}" for i, h in enumerate(reader[0])]
    rows = []
    for r in reader[1:]:
        cells = [str(c).strip()[:5000] if c is not None else "" for c in r]
        cells += [""] * (len(headers) - len(cells))
        rows.append(cells[:len(headers)])
    if not rows:
        raise ValidationError("The file has a heading row but no data rows.")
    if len(rows) > MAX_ROWS:
        raise ValidationError(f"The file has more than {MAX_ROWS} rows; split it into smaller files.")
    return headers, rows


def _cell(c):
    if c is None:
        return ""
    if isinstance(c, datetime.datetime):
        return c.strftime("%Y-%m-%d %H:%M") if (c.hour or c.minute) else c.strftime("%Y-%m-%d")
    if isinstance(c, datetime.date):
        return c.isoformat()
    if isinstance(c, datetime.time):
        return c.strftime("%H:%M")
    if isinstance(c, float) and c.is_integer():
        return str(int(c))
    return str(c).strip()


def guess_mapping(kind, headers):
    mapping = {}
    normalized = {}
    for i, h in enumerate(headers):
        normalized.setdefault(norm(h.replace("*", "")), i)
    used = set()
    for f in KINDS[kind]["fields"]:
        for syn in [f.name.replace("_", " "), f.label] + f.synonyms:
            i = normalized.get(norm(syn))
            if i is not None and i not in used:
                mapping[f.name] = i
                used.add(i)
                break
    return mapping


def missing_required(kind, mapping):
    return [f.label for f in KINDS[kind]["fields"] if f.required and f.name not in mapping]


# ------------------------------------------------------------------ 2. checking each row
def _read_row(batch, raw, ctx):
    data, errors, warnings = {}, [], []
    for f in KINDS[batch.kind]["fields"]:
        col = batch.mapping.get(f.name)
        cell = raw[col] if col is not None and col < len(raw) else ""
        value, err = read_cell(f, cell, ctx)
        if err:
            errors.append(err)
        elif value is None and f.required:
            errors.append(f"{f.label} is required")
        if f.type == "customer" and isinstance(value, tuple):
            value, note = value
            if note:
                warnings.append(f"Customer {note}: {value.code} {value.company_name}")
        data[f.name] = value
    return data, errors, warnings


def _owner(kind, data, ctx, errors):
    field = OWNER_FIELD[kind]
    if data.get(field) is None and kind != "customers":
        if ctx.default_owner is None:
            errors.append("Employee is required: the CRM does not assign imported records to an administrator "
                          "or director automatically")
        else:
            data[field] = ctx.default_owner


def analyse(batch, ctx):
    """Classify every row. Returns a list of dicts: row, action, data, messages, changes, target, keys."""
    from .models import ImportKey

    out, seen = [], {}
    for idx, raw in enumerate(batch.rows, start=2):  # row 1 = headings, as Excel numbers them
        data, errors, warnings = _read_row(batch, raw, ctx)
        if not errors:
            _owner(batch.kind, data, ctx, errors)
        if not errors:
            errors += _business_rules(batch.kind, data)
        res = {"row": idx, "action": "error", "data": data, "messages": errors + warnings, "changes": [],
               "target": None, "keys": []}
        if errors:
            out.append(res)
            continue
        handler = HANDLERS[batch.kind]
        keys = [k for k in (ref_key(data), handler.fingerprint(data)) if k]
        res["keys"] = keys
        in_file = next((k for k in keys if k in seen), None) or next(
            (k for k in handler.file_keys(data) if k in seen), None)
        if in_file:
            res.update(action="duplicate", messages=[f"Same record as Excel row {seen[in_file]} in this file"] +
                       warnings)
            out.append(res)
            continue
        for k in keys + handler.file_keys(data):
            seen[k] = idx
        reliable = handler.reliable_match(data)
        if reliable is None:
            prior = ImportKey.objects.filter(kind=batch.kind, key__in=keys).first()
            if prior is not None:
                obj = handler.model().objects.filter(pk=prior.object_id).first()
                if obj is not None:
                    reliable = (obj, "imported before with the same " + (
                        "reference" if prior.key.startswith("ref:") else "details"))
        if reliable is not None:
            obj, why = reliable
            if batch.mode == "skip":
                res.update(action="duplicate", target=obj, messages=[f"Already in the CRM: {handler.label(obj)} "
                                                                     f"({why}) — skipped"] + warnings)
            else:
                changes = handler.changes(obj, data, ctx)
                if changes:
                    res.update(action="update", target=obj, changes=changes,
                               messages=[f"Will update {handler.label(obj)} ({why})"] + warnings)
                else:
                    res.update(action="duplicate", target=obj, messages=[f"Already in the CRM and identical: "
                                                                         f"{handler.label(obj)}"] + warnings)
            out.append(res)
            continue
        possible = handler.possible_duplicate(data)
        if possible:
            res.update(action="duplicate", messages=[f"Possible duplicate of {possible} — skipped (not a reliable "
                                                     f"match, so it is never updated)"] + warnings)
        elif batch.mode == "update_only":
            res.update(action="duplicate", messages=["No existing record matches — skipped (update-only mode)"] +
                       warnings)
        else:
            res.update(action="new", messages=warnings)
        out.append(res)
    return out


def _business_rules(kind, data):
    errors = []
    today = timezone.localdate()
    if kind == "leads" and data.get("created_date") and data["created_date"] > today:
        errors.append("Date received cannot be in the future")
    if kind == "visits":
        if data.get("status") is None:
            data["status"] = "completed" if data["visit_date"] <= today else "scheduled"
        if data["status"] == "completed" and data["visit_date"] > today:
            errors.append("A visit in the future cannot be completed")
        if data["status"] == "completed" and not data.get("meeting_notes"):
            errors.append("Meeting outcome / notes are required for a completed visit")
        start = data.get("start_time") or datetime.time(9, 0)
        end = data.get("end_time") or (datetime.datetime.combine(today, start) + datetime.timedelta(hours=1)).time()
        if end <= start:
            errors.append("End time must be after the start time")
        data["start_time"], data["end_time"] = start, end
        lead, opp = data.get("lead"), data.get("opportunity")
        if lead and lead.customer_id and lead.customer_id != data["customer"].pk:
            errors.append(f"Lead {lead.code} belongs to another customer")
        if opp and opp.customer_id != data["customer"].pk:
            errors.append(f"Opportunity {opp.code} belongs to another customer")
    if kind == "tasks":
        data["category"] = data.get("category") or "follow_up"
        data["status"] = data.get("status") or "open"
        if data.get("completed_date") and data["completed_date"] > today:
            errors.append("Completed date cannot be in the future")
        lead, cust = data.get("lead"), data.get("customer")
        if lead and cust and lead.customer_id and lead.customer_id != cust.pk:
            errors.append(f"Lead {lead.code} belongs to another customer")
    if kind == "opportunities":
        st = data["stage"]
        if st.is_won or st.is_lost:
            if not data.get("closed_date"):
                errors.append("Closing date is required for Won and Lost deals")
            elif data["closed_date"] > today:
                errors.append("Closing date cannot be in the future")
        if st.is_lost and not data.get("lost_reason"):
            errors.append("Lost reason is required for Lost deals")
        if st.is_won and not data.get("estimated_value"):
            errors.append("A Won deal needs a value greater than zero")
    return errors


# ------------------------------------------------------------------ 3. preview and 4. import
def preview(batch, user):
    ctx = Context(user, batch.kind)
    rows = analyse(batch, ctx)
    batch.results = [{"row": r["row"], "status": r["action"], "messages": r["messages"],
                      "changes": [[a, str(b), str(c)] for a, b, c in r["changes"]]} for r in rows]
    batch.summary = _summary(rows)
    batch.status = batch.Status.PREVIEW
    batch.save()
    return rows


def _summary(rows):
    count = {a: sum(1 for r in rows if r["action"] == a) for a in ("new", "update", "duplicate", "error")}
    return {"total": len(rows), "new": count["new"], "updates": count["update"], "duplicates": count["duplicate"],
            "errors": count["error"], "ok": count["new"] + count["update"]}


def run_import(batch, user):
    from .models import ImportKey

    if batch.status != batch.Status.PREVIEW:
        raise ValidationError("Preview the import before running it.")
    ctx = Context(user, batch.kind)
    handler = HANDLERS[batch.kind]
    created = updated = 0
    with transaction.atomic():  # all or nothing
        rows = analyse(batch, ctx)  # checked again: data may have changed since the preview
        for r in rows:
            if r["action"] == "new":
                obj = handler.create(r["data"], ctx, batch)
                created += 1
            elif r["action"] == "update":
                obj = r["target"]
                handler.update(obj, r["data"], ctx, batch)
                updated += 1
            else:
                continue
            for k in r["keys"]:
                ImportKey.objects.get_or_create(kind=batch.kind, key=k, defaults={"object_id": obj.pk,
                                                                                  "batch": batch})
        batch.results = [{"row": r["row"], "status": {"new": "created", "update": "updated"}.get(r["action"],
                                                                                                 r["action"]),
                          "messages": r["messages"], "changes": [[a, str(b), str(c)] for a, b, c in r["changes"]]}
                         for r in rows]
        batch.summary = {**_summary(rows), "created": created, "updated": updated}
        # keep only the rows that were NOT imported (for the problem report); drop the rest of the file
        keep = {r["row"] for r in rows if r["action"] in ("error", "duplicate")}
        batch.rows = [raw if (i + 2) in keep else [] for i, raw in enumerate(batch.rows)]
        batch.status = batch.Status.IMPORTED
        batch.imported_at = timezone.now()
        batch.save()
        log_action("other", batch, message=f"Imported {batch.get_kind_display().lower()} from {batch.file_name}: "
                                           f"{created} created, {updated} updated, "
                                           f"{batch.summary['duplicates']} duplicates and {batch.summary['errors']} "
                                           f"errors skipped")
    return created, updated


def problem_rows(batch):
    out = []
    for res, raw in zip(batch.results, batch.rows):
        if res["status"] in ("error", "duplicate"):
            out.append([res["row"], "Error" if res["status"] == "error" else "Duplicate (skipped)",
                        "; ".join(res["messages"])] + list(raw))
    return out


# ------------------------------------------------------------------ helpers shared by the handlers
def _at(d, hour=9):
    """A date from the file as an aware datetime at 09:00 Dubai time (keeps the real day in reports)."""
    return timezone.make_aware(datetime.datetime.combine(d, datetime.time(hour, 0)))


def _set_dates(obj, **fields):
    type(obj).objects.filter(pk=obj.pk).update(**fields)
    for k, v in fields.items():
        setattr(obj, k, v)


def _diff(obj, data, names, labels):
    changes = []
    for n in names:
        new = data.get(n)
        if new in (None, ""):
            continue  # an empty cell never blanks an existing value
        old = getattr(obj, n)
        if old != new:
            changes.append((labels.get(n, n), old if old not in (None, "") else "—", new))
    return changes


def _apply(obj, data, names):
    for n in names:
        if data.get(n) not in (None, ""):
            setattr(obj, n, data[n])


def _labels(kind):
    return {f.name: f.label for f in KINDS[kind]["fields"]}


class CustomerHandler:
    names = ["category", "industry", "status", "contact_person", "designation", "phone", "alt_phone", "email",
             "website", "billing_address", "city", "country", "trn", "source", "assigned_engineer", "notes"]

    @staticmethod
    def model():
        from customers.models import Customer
        return Customer

    def label(self, c):
        return f"{c.code} {c.company_name}"

    def fingerprint(self, d):
        return None  # customers are matched by code / TRN / email+phone / reference instead

    def file_keys(self, d):
        from customers.models import normalize_name, normalize_phone
        keys = [f"name:{normalize_name(d['company_name'])}"]
        if d.get("trn"):
            keys.append(f"trn:{d['trn'].replace(' ', '').upper()}")
        if d.get("email"):
            keys.append(f"email:{d['email']}")
        if d.get("phone") and len(normalize_phone(d["phone"])) >= 7:
            keys.append(f"phone:{normalize_phone(d['phone'])}")
        return keys

    def reliable_match(self, d):
        from customers.models import Customer, normalize_phone
        if d.get("crm_code"):
            c = Customer.objects.filter(code__iexact=d["crm_code"]).first()
            return (c, "same customer code") if c else None
        if d.get("trn"):
            c = Customer.objects.filter(trn=d["trn"].replace(" ", "").upper()).first()
            if c:
                return c, "same TRN"
        if d.get("email") and d.get("phone"):
            for c in Customer.objects.filter(email__iexact=d["email"]):
                if normalize_phone(c.phone) == normalize_phone(d["phone"]):
                    return c, "same email and phone"
        return None

    def possible_duplicate(self, d):
        from customers.services import find_duplicates
        dups = find_duplicates(company_name=d["company_name"], email=d.get("email"), phone=d.get("phone"),
                               trn=d.get("trn"))
        return f"{dups[0]['customer']} ({', '.join(dups[0]['reasons'])})" if dups else None

    def changes(self, obj, d, ctx):
        return _diff(obj, d, self.names, _labels("customers"))

    def create(self, d, ctx, batch):
        from customers.models import Customer
        c = Customer(company_name=d["company_name"], status=d.get("status") or "prospect")
        _apply(c, d, self.names)
        if c.trn:
            c.trn = c.trn.replace(" ", "").upper()
        c.created_by = d.get("assigned_engineer") or ctx.user
        c.save()
        record(c, "created", f"Imported from {batch.file_name} by {ctx.user.get_full_name() or ctx.user.username}")
        return c

    def update(self, obj, d, ctx, batch):
        _apply(obj, d, self.names)
        obj.save()
        record(obj, "updated", f"Updated from import {batch.file_name}")


class LeadHandler:
    names = ["title", "company_name", "contact_person", "designation", "email", "phone", "city", "source",
             "interest", "estimated_value", "priority", "assigned_engineer", "next_follow_up", "qualification_notes"]

    @staticmethod
    def model():
        from leads.models import Lead
        return Lead

    def label(self, x):
        return f"{x.code} {x.title}"

    def fingerprint(self, d):
        return fingerprint("leads", d["company_name"], d["title"], d.get("created_date") or "")

    def file_keys(self, d):
        return [f"lead:{norm(d['company_name'])}|{norm(d['title'])}"]

    def reliable_match(self, d):
        from leads.models import Lead
        if d.get("crm_code"):
            x = Lead.objects.filter(code__iexact=d["crm_code"]).first()
            return (x, "same lead code") if x else None
        return None

    def possible_duplicate(self, d):
        from leads.services import find_lead_duplicates
        dups = find_lead_duplicates(company_name=d["company_name"], email=d.get("email"), phone=d.get("phone"))
        return f"open lead {dups[0].code} {dups[0].title}" if dups else None

    def changes(self, obj, d, ctx):
        ch = _diff(obj, d, self.names, _labels("leads"))
        if d.get("stage") and d["stage"] != obj.stage:
            ch.append(("Stage", obj.stage, d["stage"]))
        return ch

    def create(self, d, ctx, batch):
        from leads.models import Lead, LeadStage
        from leads.services import record_assignment
        lead = Lead(stage=d.get("stage") or LeadStage.first_open())
        _apply(lead, d, self.names)
        lead.created_by = d["assigned_engineer"]  # counted for the employee, not the importer
        lead.save()
        if d.get("created_date"):
            _set_dates(lead, created_at=_at(d["created_date"]), stage_changed_at=_at(d["created_date"]))
        record(lead, "created", f"Imported from {batch.file_name} by {ctx.user.get_full_name() or ctx.user.username}")
        record_assignment(lead, note="Imported")
        return lead

    def update(self, obj, d, ctx, batch):
        from leads.services import change_stage
        _apply(obj, d, self.names)
        obj.save()
        if d.get("stage") and d["stage"] != obj.stage:
            change_stage(obj, d["stage"], note=f"Import {batch.file_name}")
        record(obj, "updated", f"Updated from import {batch.file_name}")


class VisitHandler:
    names = ["visit_date", "start_time", "end_time", "location", "purpose", "status", "assigned_to", "lead",
             "opportunity", "meeting_notes", "next_follow_up_date"]

    @staticmethod
    def model():
        from visits.models import Visit
        return Visit

    def label(self, v):
        return f"{v.code} {v.customer.company_name} {v.visit_date:%d-%b-%Y}"

    def fingerprint(self, d):
        return fingerprint("visits", d["customer"].pk, d["visit_date"], d["start_time"].strftime("%H:%M"),
                           d["assigned_to"].pk)

    def file_keys(self, d):
        return []

    def reliable_match(self, d):
        from visits.models import Visit
        if d.get("crm_code"):
            v = Visit.objects.filter(code__iexact=d["crm_code"]).first()
            return (v, "same visit code") if v else None
        return None

    def possible_duplicate(self, d):
        from visits.models import Visit
        v = Visit.objects.filter(customer=d["customer"], visit_date=d["visit_date"], start_time=d["start_time"],
                                 assigned_to=d["assigned_to"]).first()
        return f"visit {v.code}" if v else None

    def changes(self, obj, d, ctx):
        return _diff(obj, d, self.names, _labels("visits"))

    def create(self, d, ctx, batch):
        from visits.models import Visit
        v = Visit(customer=d["customer"], location=d.get("location") or d["customer"].city or "Customer site")
        _apply(v, d, self.names)
        v.created_by = d["assigned_to"]
        if v.status == "completed":
            v.completed_by = d["assigned_to"]
            v.completed_at = _at(d["visit_date"], d["end_time"].hour)
        v.save()
        record(v, "created", f"Imported from {batch.file_name} by {ctx.user.get_full_name() or ctx.user.username}",
               customer=v.customer)
        return v

    def update(self, obj, d, ctx, batch):
        _apply(obj, d, self.names)
        if obj.status == "completed" and not obj.completed_at:
            obj.completed_by, obj.completed_at = obj.assigned_to, _at(obj.visit_date, obj.end_time.hour)
        obj.save()
        record(obj, "updated", f"Updated from import {batch.file_name}", customer=obj.customer)


class TaskHandler:
    names = ["title", "due_date", "category", "priority", "status", "assigned_to", "customer", "lead", "description"]

    @staticmethod
    def model():
        from tasks.models import Task
        return Task

    def label(self, t):
        return f"{t.code} {t.title}"

    def fingerprint(self, d):
        return fingerprint("tasks", d["title"], d["due_date"], d["assigned_to"].pk,
                           d["customer"].pk if d.get("customer") else "")

    def file_keys(self, d):
        return []

    def reliable_match(self, d):
        from tasks.models import Task
        if d.get("crm_code"):
            t = Task.objects.filter(code__iexact=d["crm_code"]).first()
            return (t, "same task code") if t else None
        return None

    def possible_duplicate(self, d):
        from tasks.models import Task
        t = Task.objects.filter(title__iexact=d["title"], due_date=d["due_date"], assigned_to=d["assigned_to"]).first()
        return f"task {t.code}" if t else None

    def changes(self, obj, d, ctx):
        return _diff(obj, d, self.names, _labels("tasks"))

    def create(self, d, ctx, batch):
        from tasks.models import Task
        t = Task()
        _apply(t, d, self.names)
        t.created_by = d["assigned_to"]
        t.save()
        if t.status == "completed":
            _set_dates(t, completed_at=_at(d.get("completed_date") or d["due_date"], 17))
        record(t, "created", f"Imported from {batch.file_name} by {ctx.user.get_full_name() or ctx.user.username}",
               customer=t.customer)
        return t

    def update(self, obj, d, ctx, batch):
        was_done = obj.status == "completed"
        _apply(obj, d, self.names)
        obj.save()
        if obj.status == "completed" and (not was_done or d.get("completed_date")):
            _set_dates(obj, completed_at=_at(d.get("completed_date") or obj.due_date, 17))
        record(obj, "updated", f"Updated from import {batch.file_name}", customer=obj.customer)


class OpportunityHandler:
    names = ["title", "estimated_value", "expected_close_date", "lost_reason", "assigned_engineer", "source", "notes"]

    @staticmethod
    def model():
        from sales.models import Opportunity
        return Opportunity

    def label(self, o):
        return f"{o.code} {o.title}"

    def fingerprint(self, d):
        return fingerprint("opportunities", d["customer"].pk, d["title"], d["estimated_value"],
                           d.get("closed_date") or "")

    def file_keys(self, d):
        return [f"opp:{d['customer'].pk}|{norm(d['title'])}"]

    def reliable_match(self, d):
        from sales.models import Opportunity
        if d.get("crm_code"):
            o = Opportunity.objects.filter(code__iexact=d["crm_code"]).first()
            return (o, "same opportunity code") if o else None
        return None

    def possible_duplicate(self, d):
        from sales.models import Opportunity
        o = Opportunity.objects.filter(customer=d["customer"], title__iexact=d["title"]).first()
        return f"opportunity {o.code} {o.title}" if o else None

    def changes(self, obj, d, ctx):
        ch = _diff(obj, d, self.names, _labels("opportunities"))
        if d["stage"] != obj.stage:
            ch.append(("Status / stage", obj.stage, d["stage"]))
        if d.get("closed_date") and (obj.closed_at is None or timezone.localtime(obj.closed_at).date() !=
                                     d["closed_date"]):
            ch.append(("Closing date", timezone.localtime(obj.closed_at).date() if obj.closed_at else "—",
                       d["closed_date"]))
        return ch

    def create(self, d, ctx, batch):
        from sales.models import Opportunity, OpportunityStageHistory
        st = d["stage"]
        o = Opportunity(customer=d["customer"], stage=st)
        _apply(o, d, self.names)
        o.probability = 100 if st.is_won else 0 if st.is_lost else (d.get("probability")
                                                                    if d.get("probability") is not None
                                                                    else st.default_probability)
        o.created_by = d["assigned_engineer"]
        o.save()
        if st.is_won or st.is_lost:
            when = _at(d["closed_date"])
            _set_dates(o, closed_at=when, stage_changed_at=when, created_at=min(when, o.created_at))
        OpportunityStageHistory.objects.create(opportunity=o, from_stage=None, to_stage=st, probability=o.probability,
                                               estimated_value=o.estimated_value, changed_by=d["assigned_engineer"],
                                               changed_at=o.closed_at or timezone.now(),
                                               note=f"Imported from {batch.file_name}")
        record(o, "created", f"Imported from {batch.file_name} by {ctx.user.get_full_name() or ctx.user.username}"
                             + (f" — {st.name} on {d['closed_date']:%d-%b-%Y}" if d.get("closed_date") else ""),
               customer=o.customer)
        return o

    def update(self, obj, d, ctx, batch):
        from sales.models import OpportunityStageHistory
        _apply(obj, d, self.names)
        st = d["stage"]
        if st != obj.stage:
            OpportunityStageHistory.objects.create(opportunity=obj, from_stage=obj.stage, to_stage=st,
                                                   probability=100 if st.is_won else 0 if st.is_lost
                                                   else obj.probability, estimated_value=obj.estimated_value,
                                                   changed_by=obj.assigned_engineer,
                                                   changed_at=_at(d["closed_date"]) if d.get("closed_date")
                                                   else timezone.now(), note=f"Import {batch.file_name}")
            obj.stage = st
            obj.probability = 100 if st.is_won else 0 if st.is_lost else obj.probability
        obj.save()
        if (st.is_won or st.is_lost) and d.get("closed_date"):
            _set_dates(obj, closed_at=_at(d["closed_date"]), stage_changed_at=_at(d["closed_date"]))
        record(obj, "updated", f"Updated from import {batch.file_name}", customer=obj.customer)


HANDLERS = {"customers": CustomerHandler(), "leads": LeadHandler(), "visits": VisitHandler(), "tasks": TaskHandler(),
            "opportunities": OpportunityHandler()}


# ------------------------------------------------------------------ templates
def template_workbook(kind):
    """An .xlsx template: the Data sheet (headings + 2 fictional example rows) and an Instructions sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from core.models import LookupValue
    from leads.models import LeadStage
    from sales.models import OpportunityStage

    spec = KINDS[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    fields = spec["fields"]
    ws.append([f.label + (" *" if f.required else "") for f in fields])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E6EBA0")
    for row in EXAMPLES.get(kind, []):
        ws.append([row.get(f.name, "") for f in fields])
    for i, f in enumerate(fields, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, min(40, len(f.label) + 6))
    ins = wb.create_sheet("Instructions")
    ins.append([f"{spec['label']} import template — Technoflow CRM"])
    ins["A1"].font = Font(bold=True, size=13)
    ins.append(["Fill the Data sheet (one row per record; delete the two example rows). Columns marked * are "
                "required. Dates: YYYY-MM-DD or DD/MM/YYYY. Times: HH:MM. Numbers without currency signs."])
    ins.append(["Nothing is saved until you check the preview and click Import. Duplicates are skipped; existing "
                "records are only updated if you choose an update option and the match is reliable "
                "(CRM code, your Reference, TRN)."])
    ins.append([])
    ins.append(["Column", "Required", "What to enter", "Allowed values"])
    for c in ins[5]:
        c.font = Font(bold=True)
    for f in fields:
        allowed = ""
        if f.type == "choice":
            allowed = ", ".join(sorted({v.replace("_", " ").title() for v in f.choices.values()}))
        elif f.type.startswith("lookup:"):
            allowed = ", ".join(LookupValue.objects.filter(category=f.type.split(":", 1)[1], is_active=True)
                                .values_list("value", flat=True))
        elif f.type == "lead_stage":
            allowed = ", ".join(LeadStage.objects.filter(is_active=True, is_won=False, is_lost=False)
                                .values_list("name", flat=True))
        elif f.type == "opp_stage":
            allowed = "Won, Lost, " + ", ".join(OpportunityStage.objects.filter(
                is_active=True, is_won=False, is_lost=False).values_list("name", flat=True))
        elif f.type == "date":
            allowed = "2026-03-15 or 15/03/2026"
        ins.append([f.label, "Yes" if f.required else "No", f.help, allowed])
    for col, width in zip("ABCD", (30, 10, 70, 60)):
        ins.column_dimensions[col].width = width
    for row in ins.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


EXAMPLES = {  # fictional examples only
    "customers": [{"company_name": "Example Contracting LLC", "category": "", "status": "Active",
                   "contact_person": "Mr. Sample Person", "phone": "+971 4 555 0101",
                   "email": "info@example-contracting.example", "city": "Dubai", "country": "United Arab Emirates",
                   "reference": "CUST-001"},
                  {"company_name": "Sample Facilities Management", "status": "Prospect", "city": "Abu Dhabi",
                   "email": "projects@sample-fm.example", "reference": "CUST-002"}],
    "leads": [{"title": "Chiller replacement - Tower A", "company_name": "Example Contracting LLC",
               "contact_person": "Mr. Sample Person", "email": "buyer@example-contracting.example",
               "estimated_value": "150000", "priority": "High", "stage": "New", "created_date": "2026-03-15",
               "reference": "LEAD-001"},
              {"title": "Pump room upgrade", "company_name": "Sample Facilities Management",
               "estimated_value": "64000", "priority": "Medium", "stage": "Contacted", "created_date": "2026-03-16",
               "reference": "LEAD-002"}],
    "visits": [{"customer": "CUS-00001", "visit_date": "2026-03-20", "start_time": "10:00", "end_time": "11:00",
                "status": "Completed", "meeting_notes": "Discussed requirements; offer to follow.",
                "next_follow_up_date": "2026-03-27", "reference": "VISIT-001"},
               {"customer": "CUS-00002", "visit_date": "2026-04-02", "start_time": "14:00", "status": "Scheduled",
                "reference": "VISIT-002"}],
    "tasks": [{"title": "Call back about the offer", "due_date": "2026-03-25", "category": "Follow-up",
               "priority": "High", "status": "Open", "customer": "CUS-00001", "reference": "TASK-001"},
              {"title": "Send company profile", "due_date": "2026-03-18", "category": "Follow-up",
               "status": "Completed", "completed_date": "2026-03-18", "reference": "TASK-002"}],
    "opportunities": [{"title": "Supply of chilled water valves", "customer": "CUS-00001", "stage": "Won",
                       "estimated_value": "250000", "closed_date": "2026-02-28", "reference": "SALE-2026-001"},
                      {"title": "Pump replacement", "customer": "CUS-00002", "stage": "Negotiation",
                       "estimated_value": "80000", "expected_close_date": "2026-05-31", "reference": "DEAL-002"}],
}
